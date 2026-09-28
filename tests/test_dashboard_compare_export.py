# -*- coding: utf-8 -*-
"""
Сравнение периодов и выгрузки Excel/PDF дашборда — числа ровно как на экране (2026-09-28).

Ошибки до правки:
1. POST /api/comparison/periods был заглушкой: всегда {comparison: {}, insights: []}.
2. Выгрузки Excel/PDF брали факт отдельным запросом в «сырых» единицах расчёта:
   наценка «Факт» дробью (2.24) против плана в процентах (200), активность кранов в
   факте — всегда 0 (calculate_metrics её не считает), пустой период — 500.

Что проверяется:
- факт выгрузки и обоих периодов сравнения = ответ /api/dashboard-analytics за те
  же даты (наценка в процентах, tapActivity из журнала кранов);
- формулы сравнения: Δ = период 1 − период 2, Δ% = Δ / период 2 × 100 (база 0 —
  null), у процентных метрик Δ в п.п., «лучше/хуже» (списания — бюджет), топ-3 по
  |Δ%| как на вкладке «Сравнение», выводы без эмодзи;
- строки выгрузки (ExportManager.dashboard_rows): план и факт в одних единицах,
  метрика без плана — пусто/«—», план активности кранов по умолчанию 100 %,
  бюджетная метрика — светофор от зеркального процента;
- Excel: наценка «Факт» в процентах, активность кранов из журнала, пустая ячейка
  плана; PDF (HTML без reportlab): те же числа; пустой период — 200 с нулями;
- проверки ввода — 400, сбой iiko — 500.

Сеть не трогается: OLAP и журнал кранов подменены.
Запуск: py -3 -m pytest -q tests/test_dashboard_compare_export.py
"""

import os
import re
import sys
import unittest
from io import BytesIO
from unittest.mock import patch

from flask import Flask

REPO_ROOT = os.path.join(os.path.dirname(__file__), "..")
sys.path.insert(0, REPO_ROOT)

from core.comparison_calculator import COMPARISON_METRICS, ComparisonCalculator  # noqa: E402
from core.export_manager import DASHBOARD_EXPORT_METRICS, ExportManager  # noqa: E402
from extensions import DASHBOARD_OLAP_CACHE  # noqa: E402
from routes.dashboard import dashboard_bp  # noqa: E402


def olap_row(order_id, revenue, cost, category, discount=0.0, card=''):
    """Строка единого OLAP-отчёта дашборда с полями, которые читает расчёт."""
    return {
        'UniqOrderId.Id': order_id, 'DishDiscountSumInt': revenue,
        'ProductCostBase.ProductCost': cost, 'DishGroup.TopParent': category,
        'DiscountSum': discount, 'Delivery.CustomerCardNumber': card,
    }


# Неделя 1 (база, «было») и неделя 2 (сравниваемая): выручка 1000 -> 1500, наценка
# растёт, списания растут (бюджет — хуже), кухни в базе нет (Δ% кухни = null).
WEEK_BASE = ('2026-03-02', '2026-03-08')
WEEK_NEW = ('2026-03-09', '2026-03-15')
ROWS = {
    WEEK_BASE[0]: [
        olap_row('A', 600.0, 300.0, 'Напитки Розлив', discount=20.0, card='7900'),
        olap_row('B', 400.0, 250.0, 'Напитки Фасовка'),
    ],
    WEEK_NEW[0]: [
        olap_row('C', 900.0, 300.0, 'Напитки Розлив', discount=50.0, card='7901'),
        olap_row('D', 300.0, 150.0, 'Напитки Фасовка'),
        olap_row('E', 300.0, 100.0, 'ЕДА'),
    ],
}
TAP_ACTIVITY = {WEEK_BASE[0]: 80.0, WEEK_NEW[0]: 90.0}


class RangeOlap:
    """iiko: строки по дате начала периода; незнакомый период — пустой (закрытый день)."""

    def __init__(self):
        self.api = type('Api', (), {'base_url': 'https://example.invalid/resto/api'})()

    def connect(self):
        return True

    def disconnect(self):
        return None

    def get_all_sales_report(self, date_from, date_to, bar_name=None):
        return {'data': [dict(row) for row in ROWS.get(date_from, [])]}


class BrokenOlap(RangeOlap):
    def get_all_sales_report(self, date_from, date_to, bar_name=None):
        return None


def tap_activity(bar_id, date_from, date_to):
    return TAP_ACTIVITY.get(date_from, 0.0)


def key(period):
    return period[0] + '_' + period[1]


@patch('routes.dashboard.taps_manager.calculate_tap_activity_for_period', side_effect=tap_activity)
@patch('routes.dashboard.OlapReports', RangeOlap)
class PageParity(unittest.TestCase):
    """Числа сравнения и выгрузки = числа экрана за те же даты."""

    def setUp(self):
        DASHBOARD_OLAP_CACHE.clear()
        app = Flask(__name__)
        app.register_blueprint(dashboard_bp)
        self.client = app.test_client()

    def page(self, period, bar='all'):
        response = self.client.post('/api/dashboard-analytics', json={
            'bar': bar, 'date_from': period[0], 'date_to': period[1]})
        self.assertEqual(200, response.status_code, response.get_data(as_text=True))
        return response.get_json()

    def compare(self, **body):
        payload = {'venue_key': 'all', 'period1_key': key(WEEK_NEW),
                   'period2_key': key(WEEK_BASE)}
        payload.update(body)
        return self.client.post('/api/comparison/periods', json=payload)

    # ------------------------------------------------------------ сравнение

    def test_compare_is_real_and_matches_page(self, *_mocks):
        """Регрессия: раньше заглушка отвечала пустым comparison."""
        response = self.compare()
        self.assertEqual(200, response.status_code, response.get_data(as_text=True))
        body = response.get_json()
        self.assertEqual('period2', body['base'])
        self.assertEqual(20, len(body['comparison']))
        page_new, page_base = self.page(WEEK_NEW), self.page(WEEK_BASE)
        for metric_key, _label, _unit in COMPARISON_METRICS:
            row = body['comparison'][metric_key]
            self.assertAlmostEqual(page_new[metric_key], row['period1'], places=6, msg=metric_key)
            self.assertAlmostEqual(page_base[metric_key], row['period2'], places=6, msg=metric_key)
            self.assertEqual(page_new[metric_key], body['period1']['metrics'][metric_key])
        self.assertEqual({'key': key(WEEK_NEW), 'date_from': WEEK_NEW[0],
                          'date_to': WEEK_NEW[1]},
                         {k: body['period1'][k] for k in ('key', 'date_from', 'date_to')})

    def test_compare_formulas(self, *_mocks):
        body = self.compare().get_json()
        revenue = body['comparison']['revenue']
        self.assertEqual((1500.0, 1000.0, 500.0, 50.0),
                         (revenue['period1'], revenue['period2'], revenue['diff'],
                          revenue['diff_percent']))
        self.assertEqual(('up', True, '₽'), (revenue['trend'], revenue['better'],
                                             revenue['diff_unit']))
        markup = body['comparison']['markupPercent']
        self.assertGreater(markup['period1'], 100, 'наценка в процентах, не дробью')
        self.assertEqual('п.п.', markup['diff_unit'])
        self.assertAlmostEqual(markup['period1'] - markup['period2'], markup['diff'], places=2)
        kitchen = body['comparison']['revenueKitchen']
        self.assertIsNone(kitchen['diff_percent'], 'база 0 — процент изменения не определён')
        writeoffs = body['comparison']['loyaltyWriteoffs']
        self.assertEqual(('up', False, True), (writeoffs['trend'], writeoffs['better'],
                                               writeoffs['budget']))
        taps = body['comparison']['tapActivity']
        self.assertEqual((90.0, 80.0, 10.0), (taps['period1'], taps['period2'], taps['diff']))

    def test_top_changes_and_insights(self, *_mocks):
        body = self.compare().get_json()
        top = body['top_changes']
        self.assertLessEqual(len(top), 3)
        percents = [abs(row['diff_percent']) for row in top]
        self.assertEqual(sorted(percents, reverse=True), percents)
        self.assertNotIn('revenueKitchen', [row['metric'] for row in top], 'база 0 — не в топе')
        self.assertEqual(len(top), len(body['insights']))
        emoji = re.compile('[\U0001F000-\U0001FAFF\u2600-\u27BF\u2B00-\u2BFF]')
        for line in body['insights'] + [body['formula']]:
            self.assertFalse(emoji.search(line), line)
        self.assertIn('против', body['insights'][0])

    def test_compare_bad_input_is_400(self, *_mocks):
        for body, why in (({'period1_key': '2026-03'}, 'месячный ключ'),
                          ({'period2_key': '2026-03-15_2026-03-09'}, 'начало позже конца'),
                          ({'period1_key': '2026-02-30_2026-03-01'}, 'несуществующая дата'),
                          ({'venue_key': 'Лиговский'}, 'русское имя бара'),
                          ({'period1_key': None}, 'нет периода')):
            response = self.compare(**body)
            self.assertEqual(400, response.status_code, why)
        self.assertEqual(400, self.client.post('/api/comparison/periods', json=[1]).status_code)

    def test_compare_empty_venue_means_network(self, *_mocks):
        self.assertEqual('all', self.compare(venue_key='').get_json()['venue_key'])

    # ------------------------------------------------------------ выгрузки

    def excel_rows(self, period, bar='all'):
        from openpyxl import load_workbook
        response = self.client.post('/api/export/excel', json={
            'bar': bar, 'date_from': period[0], 'date_to': period[1]})
        # Тело — двоичный xlsx: в сообщение идут только первые байты.
        self.assertEqual(200, response.status_code, response.data[:300])
        self.assertIn('spreadsheetml', response.headers['Content-Type'])
        sheet = load_workbook(BytesIO(response.data)).active
        rows = {}
        for label, plan, fact in sheet.iter_rows(min_row=5, max_col=3, values_only=True):
            if label:
                rows[label] = (plan, fact)
        return rows

    def test_excel_fact_matches_page_units(self, *_mocks):
        """Регрессия: наценка «Факт» была дробью, активность кранов — 0."""
        page = self.page(WEEK_NEW)
        rows = self.excel_rows(WEEK_NEW)
        self.assertEqual(20, len(rows))
        self.assertAlmostEqual(round(page['markupPercent'], 2), rows['Наценка (%)'][1], places=2)
        self.assertGreater(rows['Наценка (%)'][1], 100)
        self.assertEqual(90.0, rows['Активность кранов (%)'][1])
        self.assertEqual(round(page['revenue'], 2), rows['Выручка (₽)'][1])
        # Наценка плана и факта — обе в процентах (план из data/plansdashboard.json,
        # если он есть за эти даты; иначе пустая ячейка, а не 0).
        plan = rows['Наценка (%)'][0]
        self.assertTrue(plan is None or plan > 10, plan)
        self.assertIsNone(rows['Чеки с картой (шт)'][0], 'плана нет — пустая ячейка')

    def test_excel_empty_period_is_zeros_not_500(self, *_mocks):
        rows = self.excel_rows(('2026-04-01', '2026-04-01'))
        self.assertEqual(0, rows['Выручка (₽)'][1])

    def test_pdf_html_matches_page(self, *_mocks):
        page = self.page(WEEK_NEW)
        response = self.client.post('/api/export/pdf', json={
            'bar': 'all', 'date_from': WEEK_NEW[0], 'date_to': WEEK_NEW[1]})
        self.assertEqual(200, response.status_code, response.get_data(as_text=True))
        text = response.get_data(as_text=True)
        if 'text/html' in response.headers['Content-Type']:
            self.assertIn(f"{page['markupPercent']:,.2f} %", text)
            self.assertIn('90.00 %', text)

    def test_export_bad_input_is_400(self, *_mocks):
        for body in ({'bar': 'Лиговский', 'date_from': '2026-03-01', 'date_to': '2026-03-02'},
                     {'bar': 'all', 'date_from': '01.03.2026', 'date_to': '2026-03-02'},
                     {'bar': 'all', 'date_from': '2026-03-05', 'date_to': '2026-03-02'},
                     {'bar': 'all'}):
            for route in ('/api/export/excel', '/api/export/pdf'):
                self.assertEqual(400, self.client.post(route, json=body).status_code,
                                 (route, body))


@patch('routes.dashboard.taps_manager.calculate_tap_activity_for_period', return_value=0)
@patch('routes.dashboard.OlapReports', BrokenOlap)
class IikoFailure(unittest.TestCase):
    def setUp(self):
        DASHBOARD_OLAP_CACHE.clear()
        app = Flask(__name__)
        app.register_blueprint(dashboard_bp)
        self.client = app.test_client()

    def test_failure_is_500(self, *_mocks):
        body = {'bar': 'all', 'date_from': '2026-03-02', 'date_to': '2026-03-08'}
        self.assertEqual(500, self.client.post('/api/export/excel', json=body).status_code)
        self.assertEqual(500, self.client.post('/api/export/pdf', json=body).status_code)
        response = self.client.post('/api/comparison/periods', json={
            'venue_key': 'all', 'period1_key': key(WEEK_NEW), 'period2_key': key(WEEK_BASE)})
        self.assertEqual(500, response.status_code)


class ExportRows(unittest.TestCase):
    """ExportManager.dashboard_rows — правила экрана (analytics.js buildStats)."""

    def rows(self, plan, fact):
        return {row['key']: row for row in ExportManager().dashboard_rows(plan, fact)}

    def test_order_and_units(self):
        rows = ExportManager().dashboard_rows({}, {})
        self.assertEqual([k for k, _l, _u in DASHBOARD_EXPORT_METRICS], [r['key'] for r in rows])
        self.assertEqual({k for k, _l, _u in COMPARISON_METRICS}, {r['key'] for r in rows})

    def test_plan_and_fact_in_same_units(self):
        rows = self.rows({'markupPercent': 200.0, 'revenue': 1000.0},
                         {'markupPercent': 224.0, 'revenue': 1100.0})
        self.assertAlmostEqual(112.0, rows['markupPercent']['percent'])
        self.assertAlmostEqual(24.0, rows['markupPercent']['diff'])
        self.assertEqual('ok', rows['revenue']['status'])

    def test_no_plan_is_none(self):
        rows = self.rows({'revenue': 0}, {'revenue': 500.0, 'cardChecks': 10})
        for metric in ('revenue', 'cardChecks'):
            self.assertIsNone(rows[metric]['plan'], metric)
            self.assertIsNone(rows[metric]['percent'], metric)
            self.assertEqual('none', rows[metric]['status'], metric)

    def test_tap_activity_default_plan_is_100(self):
        self.assertEqual(100.0, self.rows({}, {'tapActivity': 90.0})['tapActivity']['plan'])
        self.assertEqual(100.0, self.rows({'tapActivity': 0}, {})['tapActivity']['plan'])
        self.assertEqual(80.0, self.rows({'tapActivity': 80.0}, {})['tapActivity']['plan'])

    def test_budget_metric_status_is_mirrored(self):
        over = self.rows({'loyaltyWriteoffs': 100.0}, {'loyaltyWriteoffs': 120.0})
        under = self.rows({'loyaltyWriteoffs': 100.0}, {'loyaltyWriteoffs': 80.0})
        self.assertEqual('bad', over['loyaltyWriteoffs']['status'], 'перерасход бюджета 120 %')
        self.assertEqual('ok', under['loyaltyWriteoffs']['status'], 'экономия бюджета')
        self.assertTrue(over['loyaltyWriteoffs']['budget'])


class CalculatorUnits(unittest.TestCase):
    def test_zero_base_and_budget(self):
        comparison = ComparisonCalculator().compare_periods(
            {'revenue': 100.0, 'loyaltyWriteoffs': 50.0}, {'revenue': 0, 'loyaltyWriteoffs': 40.0})
        self.assertIsNone(comparison['revenue']['diff_percent'])
        self.assertEqual(25.0, comparison['loyaltyWriteoffs']['diff_percent'])
        self.assertFalse(comparison['loyaltyWriteoffs']['better'])
        self.assertIsNone(comparison['checks']['better'], 'без изменения — ни лучше, ни хуже')


if __name__ == '__main__':
    unittest.main(verbosity=2)
