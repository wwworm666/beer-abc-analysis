"""
Тесты конструктора OLAP-отчётов (/explorer): core/olap_catalog.py, core/olap_constructor.py,
core/olap_client.py, core/olap_saved_reports.py, core/olap_export.py, routes/explorer.py.

Самозапуск: `py -3 tests/test_olap_constructor.py`; pytest:
`py -3 -m pytest -q tests/test_olap_constructor.py`.

iiko подменён tests/olap_fake_iiko.py: он группирует «позиции чеков» по-настоящему и
считает уникальные чеки по фактам, поэтому тесты ловят главную ошибку сводных —
сложенные по строкам уникальные чеки. Сеть и файлы сервиса не трогаются: сохранённые
отчёты пишутся во временный файл, запасные копии каталога на диск не пишутся
(постоянного диска /kultura в тестах нет).
"""
import io
import json
import os
import sys
import tempfile
from datetime import date

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
sys.path.insert(0, HERE)
os.environ.setdefault('SESSION_COOKIE_SECURE', '0')

import requests  # noqa: E402

from core import (olap_catalog, olap_client, olap_constructor, olap_export,  # noqa: E402
                  olap_saved_reports)
from core.olap_client import IikoError  # noqa: E402
from core.olap_constructor import ConstructorError  # noqa: E402
from olap_fake_iiko import FACTS, FakeSession  # noqa: E402

SALES_BASE = {
    'report_type': 'SALES', 'date_from': '2026-09-21', 'date_to': '2026-09-23',
}


def _fresh(**session_kwargs):
    """Чистые кэши и поддельный iiko."""
    olap_client.clear_cache()
    olap_catalog.reset_cache()
    olap_constructor.reset_presets_cache()
    FakeSession.log = []
    olap_client.set_session_factory(lambda: FakeSession(**session_kwargs))


def _posts():
    return [entry[1] for entry in FakeSession.log if entry[0] == 'POST']


def _report(**overrides):
    payload = dict(SALES_BASE)
    payload.update(overrides)
    return olap_constructor.run_report(payload)


def _expect_error(fn, *args, code=None, text=None):
    try:
        fn(*args)
    except ConstructorError as error:
        if code:
            assert error.code == code, (error.code, error.message)
        if text:
            assert text in error.message, error.message
        return error
    raise AssertionError('ожидалась ConstructorError')


# ------------------------------------------------------------------ каталог

def test_catalog_kinds_blocked_filters_and_enums():
    _fresh()
    cat = olap_catalog.get_catalog('SALES')
    assert cat.source == 'live' and len(cat.fields) == 279
    assert cat.get('DishDiscountSumInt').kind == 'sum'
    assert cat.get('DiscountSum').kind == 'sum'
    assert cat.get('UniqOrderId').kind == 'server'
    assert cat.get('DishDiscountSumInt.average').kind == 'server'
    assert cat.get('ProductCostBase.MarkUp').kind == 'server'
    assert cat.get('Store.Name').kind is None            # разрез, не показатель
    # Фильтры: список, диапазон у числовых разрезов, даты; у поля периода и у мер — нет.
    assert cat.get('Store.Name').filter_op == 'values'
    assert cat.get('DeletedWithWriteoff').filter_op == 'values'
    assert cat.get('OrderNum').filter_op == 'range'
    assert cat.get('OpenTime').filter_op == 'date_range'
    assert cat.get('OpenDate.Typed').filter_op is None
    assert 'периода' in cat.get('OpenDate.Typed').filter_note
    assert cat.get('DishDiscountSumInt').filter_op is None
    # Запреты: доли от раскладки iikoOffice и составные значения.
    assert 'Доли' in cat.get('PercentOfSummary.ByCol').blocked
    assert cat.get('Currencies.SumInCurrency').blocked
    # Коды перечислений — русские названия.
    assert cat.get('DeletedWithWriteoff').enum_label('DELETED_WITH_WRITEOFF') == 'Удалено со списанием'
    assert cat.get('PayTypes.Group').enum_label('WRITEOFF') == 'Без выручки'
    trans = olap_catalog.get_catalog('TRANSACTIONS')
    assert trans.get('TransactionType').enum_label('SESSION_WRITEOFF') == 'Реализация товаров'
    assert trans.get('StartBalance.Amount').blocked                  # тяжёлые остатки
    assert trans.get('Sum.PartOfSummaryByCol').blocked
    assert trans.get('Amount.Out').kind == 'sum' and trans.get('Product.AvgSum').kind == 'server'


def test_catalog_groups_and_json():
    _fresh()
    cat = olap_catalog.get_catalog('SALES')
    groups = cat.groups()
    assert groups[0]['title'] == olap_catalog.FAVORITES_TITLE
    assert groups[0]['ids'][0] == 'OpenDate.Typed'
    assert groups[-1]['title'] == olap_catalog.ID_GROUP_TITLE
    assert 'UniqOrderId.Id' in groups[-1]['ids']
    titles = [g['title'] for g in groups]
    assert titles.index('Время') < titles.index('Доставка')       # доставка — в конце
    full = cat.to_json()
    assert full['groups'] and full['date_field'] == 'OpenDate.Typed'
    compact = cat.to_json(compact=True)
    assert 'groups' not in compact and 'tags' not in compact['fields'][0]
    found = cat.to_json(q='списан')
    assert {f['id'] for f in found['fields']} >= {'WriteoffReason', 'WriteoffUser'}
    assert all('groups' not in x for x in [found])


def test_catalog_falls_back_to_repo_snapshot_and_errors_without_one():
    _fresh(fail_columns=True)
    cat = olap_catalog.get_catalog('TRANSACTIONS')
    assert cat.source == 'repo' and cat.live_error and len(cat.fields) == 116
    assert cat.fetched_at == '2026-09-20'            # дата снимка, а не время выкладки
    try:
        olap_catalog.get_catalog('DELIVERIES')
    except IikoError as error:
        assert 'запасной копии нет' in error.message
    else:
        raise AssertionError('у доставки нет снимка — должна быть ошибка')


def test_catalog_snapshots_live_outside_mounted_data():
    """В проде /app/data перекрыт диском сервера (файлы data/ из образа не видны), а релиз,
    меняющий data/, скрипт выкладки не выпускает — снимки каталога только в resources/."""
    for report_type, (name, stamp) in olap_catalog.REPO_SNAPSHOTS.items():
        path = os.path.join(olap_catalog.RESOURCES_DIR, name)
        assert os.path.basename(os.path.dirname(path)) == 'resources', path
        assert os.path.exists(path), path
        assert len(stamp) == 10 and stamp[4] == '-', (report_type, stamp)


# ------------------------------------------------------------------ период

def test_resolve_period_presets():
    today = date(2026, 9, 29)          # вторник
    table = {p['id']: (p['date_from'], p['date_to'])
             for p in olap_constructor.period_table(today)}
    assert table['today'] == ('2026-09-29', '2026-09-29')
    assert table['yesterday'] == ('2026-09-28', '2026-09-28')
    assert table['this_week'] == ('2026-09-28', '2026-09-29')
    assert table['last_week'] == ('2026-09-21', '2026-09-27')
    assert table['this_month'] == ('2026-09-01', '2026-09-29')
    assert table['last_month'] == ('2026-08-01', '2026-08-31')
    assert table['last_7_days'] == ('2026-09-22', '2026-09-28')
    assert table['last_30_days'] == ('2026-08-30', '2026-09-28')
    assert table['this_year'] == ('2026-01-01', '2026-09-29')
    assert table['last_year'] == ('2025-01-01', '2025-12-31')
    assert set(table) == set(olap_constructor.PERIOD_PRESETS)
    assert olap_constructor.resolve_period('last_month', date(2026, 1, 15)) == \
        (date(2025, 12, 1), date(2025, 12, 31))


# ------------------------------------------------------------------ заявка

def _spec(**payload):
    _fresh()
    data = dict(SALES_BASE, measures=['DishDiscountSumInt'], rows=['Store.Name'])
    data.update(payload)
    return olap_constructor.parse_spec(data, olap_catalog.get_catalog(data['report_type']))


def test_parse_spec_rejects_bad_requests():
    cases = [
        (dict(rows=[], columns=[]), 'no_dims'),
        (dict(measures=[]), 'no_measures'),
        (dict(rows=['Nope.Field']), 'unknown_field'),
        (dict(rows=['DishDiscountSumInt']), 'not_groupable'),
        (dict(measures=['Store.Name']), 'not_aggregatable'),
        (dict(rows=['DishAmountInt'], measures=['DishAmountInt']), 'dim_and_measure'),
        (dict(measures=['PercentOfSummary.ByCol']), 'blocked_field'),
    ]
    for payload, code in cases:
        try:
            _spec(**payload)
        except ConstructorError as error:
            assert error.code == code, (payload, error.code, error.message)
        else:
            raise AssertionError('принята заявка ' + repr(payload))
    for payload, text in [
        (dict(rows=['Store.Name', 'Store.Name']), 'дважды'),
        (dict(rows=['Store.Name'], columns=['Store.Name']), 'и в строках, и в столбцах'),
        (dict(filters=[{'field': 'OpenDate.Typed', 'op': 'date_range', 'from': '2026-09-01',
                        'to': '2026-09-02'}]), 'поле периода'),
        (dict(filters=[{'field': 'Store.Name', 'op': 'range', 'from': 1, 'to': 2}]), 'фильтр in или not_in'),
        (dict(filters=[{'field': 'Store.Name', 'op': 'in', 'values': []}]), 'непустой'),
        (dict(filters=[{'field': 'Store.Name', 'op': 'in', 'values': ['a']},
                       {'field': 'Store.Name', 'op': 'in', 'values': ['b']}]), 'второй фильтр'),
        (dict(filters=[{'field': 'OrderNum', 'op': 'range', 'from': 5, 'to': 1}]), '«до» меньше «от»'),
        (dict(filters=[{'field': 'OpenTime', 'op': 'date_range', 'from': '2026-13-01',
                        'to': '2026-09-02'}]), 'YYYY-MM-DD'),
        (dict(date_from='2026-09-10', date_to='2026-09-01'), 'раньше начала'),
        (dict(date_from=None, date_to=None), 'Нужен период'),
        (dict(period='week', date_from=None), 'Неизвестный период'),
        (dict(include_deleted='yes'), 'include_deleted'),
    ]:
        try:
            _spec(**payload)
        except ConstructorError as error:
            assert text in error.message, (payload, error.message)
        else:
            raise AssertionError('принята заявка ' + repr(payload))


def test_filters_compile_and_enum_labels_become_codes():
    spec = _spec(filters=[
        {'field': 'DeletedWithWriteoff', 'op': 'in', 'values': ['удалено со списанием', 'NOT_DELETED']},
        {'field': 'DishGroup.TopParent', 'op': 'not_in', 'values': ['ЕДА', None]},
        {'field': 'OrderNum', 'op': 'range', 'from': '10', 'to': 20.0, 'include_high': False},
        {'field': 'OpenTime', 'op': 'date_range', 'from': '2026-09-21T18:00', 'to': '2026-09-22'},
    ])
    body = olap_constructor.olap_body(spec, ['Store.Name'], False)
    f = body['filters']
    assert f['DeletedWithWriteoff'] == {'filterType': 'IncludeValues',
                                        'values': ['DELETED_WITH_WRITEOFF', 'NOT_DELETED']}
    assert f['DishGroup.TopParent'] == {'filterType': 'ExcludeValues', 'values': ['ЕДА', None]}
    assert f['OrderNum'] == {'filterType': 'Range', 'from': 10, 'to': 20, 'includeLow': True,
                             'includeHigh': False}
    assert f['OpenTime']['from'] == '2026-09-21T18:00:00.000'
    # «До 22-го» включительно — весь день 22-го, а не только его полночь.
    assert f['OpenTime']['to'] == '2026-09-22T23:59:59.999'
    # Свой фильтр по DeletedWithWriteoff есть — автоматически добавится только OrderDeleted.
    assert [g.field for g in spec.guards()] == ['OrderDeleted']
    assert f['OrderDeleted'] == {'filterType': 'IncludeValues', 'values': ['NOT_DELETED']}


def test_datetime_upper_bound_covers_whole_day_or_minute_only_when_included():
    cases = [
        ({'to': '2026-09-22T18:30'}, '2026-09-22T18:30:59.999'),          # вся минута
        ({'to': '2026-09-22T18:30:15'}, '2026-09-22T18:30:15.000'),       # секунды заданы
        ({'to': '2026-09-22', 'include_high': False}, '2026-09-22T00:00:00.000'),
    ]
    for extra, expected in cases:
        spec = _spec(filters=[dict({'field': 'OpenTime', 'op': 'date_range',
                                    'from': '2026-09-21'}, **extra)])
        body = olap_constructor.olap_body(spec, ['Store.Name'], False)
        assert body['filters']['OpenTime']['to'] == expected, (extra, body['filters']['OpenTime'])
        assert body['filters']['OpenTime']['from'] == '2026-09-21T00:00:00.000'


def test_body_period_guards_and_summary_flag():
    spec = _spec()
    body = olap_constructor.olap_body(spec, ['Store.Name'], True)
    assert body['reportType'] == 'SALES' and body['buildSummary'] == 'true'
    assert body['groupByColFields'] == []
    assert body['filters']['OpenDate.Typed'] == {
        'filterType': 'DateRange', 'periodType': 'CUSTOM', 'from': '2026-09-21', 'to': '2026-09-24'}
    assert body['filters']['DeletedWithWriteoff']['values'] == ['NOT_DELETED']
    assert body['filters']['OrderDeleted']['values'] == ['NOT_DELETED']
    spec = _spec(include_deleted=True)
    body = olap_constructor.olap_body(spec, ['Store.Name'], False)
    assert 'DeletedWithWriteoff' not in body['filters'] and body['buildSummary'] == 'false'
    trans = _spec(report_type='TRANSACTIONS', rows=['Account.Name'], measures=['Amount.Out'])
    tbody = olap_constructor.olap_body(trans, ['Account.Name'], False)
    assert set(tbody['filters']) == {'DateTime.DateTyped'}      # у проводок нет «удалённых»


def test_plan_variants():
    cat_spec = _spec
    both = cat_spec(rows=['DishGroup.TopParent'], columns=['Store.Name'],
                    measures=['DishDiscountSumInt', 'UniqOrderId'])
    cat = olap_catalog.get_catalog('SALES')
    assert olap_constructor.plan(both, cat, 'pivot') == [
        ('cells', ['Store.Name', 'DishGroup.TopParent'], True),
        ('rows', ['DishGroup.TopParent'], True)]
    assert olap_constructor.plan(both, cat, 'flat') == [
        ('main', ['DishGroup.TopParent', 'Store.Name'], True)]
    only_sum = cat_spec(rows=['DishGroup.TopParent'], columns=['Store.Name'])
    assert olap_constructor.plan(only_sum, cat, 'pivot') == [
        ('main', ['DishGroup.TopParent', 'Store.Name'], False)]
    no_cols = cat_spec(measures=['UniqOrderId'])
    assert olap_constructor.plan(no_cols, cat, 'pivot') == [('main', ['Store.Name'], True)]
    no_rows = cat_spec(rows=[], columns=['Store.Name'], measures=['UniqOrderId'])
    assert olap_constructor.plan(no_rows, cat, 'pivot') == [('main', ['Store.Name'], True)]


# ------------------------------------------------------------------ сводная

def _node(tree, *path):
    node = tree
    for key in path:
        node = next(child for child in node['ch'] if child['k'] == key)
    return node


def test_pivot_totals_do_not_double_count_checks():
    _fresh()
    out = _report(rows=['DishGroup.TopParent'], columns=['Store.Name'],
                  measures=['DishDiscountSumInt', 'UniqOrderId'], format='pivot')
    assert out['columns'] == [['Варшавская'], ['Лиговский']]
    tree = out['tree']
    # Итог столбца «Лиговский»: чеки o1, o2, o5, o6 = 4 (сложение строк дало бы 5).
    assert tree['c'] == [[550, 1], [1600, 4]]
    assert tree['t'] == [2150, 5]
    assert _node(tree, 'ЕДА')['t'] == [600, 2]           # o1 и o3
    assert _node(tree, 'Напитки Розлив')['c'] == [None, [1250, 3]]
    assert [child['k'] for child in tree['ch']][-1] is None   # пусто — в конце
    assert out['leaf_count'] == 4 and not out['warnings']
    assert [r['name'] for r in out['meta']['requests']] == ['cells', 'rows']
    assert out['meta']['totals'] == {'sum': ['DishDiscountSumInt'], 'server': ['UniqOrderId']}
    assert any('автоматически' in f for f in out['meta']['filters'])
    assert out['meta']['date_to_exclusive'] == '2026-09-24'
    assert len(_posts()) == 2


def test_pivot_nested_groups_and_sum_only_single_request():
    _fresh()
    out = _report(rows=['Store.Name', 'DishGroup.TopParent'], measures=['DishDiscountSumInt'],
                  format='pivot')
    assert len(_posts()) == 1 and _posts()[0]['buildSummary'] == 'false'
    tree = out['tree']
    assert tree['t'] == [2150]
    lig = _node(tree, 'Лиговский')
    assert lig['t'] == [1600]
    assert _node(tree, 'Лиговский', 'Напитки Розлив')['t'] == [1250]
    assert 'c' not in tree                                  # без столбцов — только «Итого»


def test_pivot_rows_only_and_columns_only_with_server_measure():
    _fresh()
    out = _report(rows=['Store.Name', 'DishGroup.TopParent'], measures=['UniqOrderId'],
                  format='pivot')
    tree = out['tree']
    assert tree['t'] == [5]
    assert _node(tree, 'Лиговский')['t'] == [4]              # итог группы — из iiko
    assert _node(tree, 'Лиговский', 'ЕДА')['t'] == [1]
    _fresh()
    out = _report(rows=[], columns=['Store.Name'], measures=['UniqOrderId', 'DishDiscountSumInt'],
                  format='pivot')
    assert out['tree']['c'] == [[1, 550], [4, 1600]]
    assert out['tree']['t'] == [5, 2150] and out['leaf_count'] == 0


def test_pivot_group_without_data_in_column_is_empty_not_missing():
    """У группы нет строк в столбце — ячейка пустая, без ложного «iiko не вернул итоги»."""
    _fresh()
    out = _report(rows=['Store.Name', 'DishGroup.TopParent'], columns=['OpenDate.Typed'],
                  measures=['UniqOrderId', 'DishDiscountSumInt'], format='pivot')
    assert out['columns'] == [['2026-09-21'], ['2026-09-22'], ['2026-09-23']]
    assert not out['warnings'], out['warnings']
    war = _node(out['tree'], 'Варшавская')
    assert war['c'][0] is None and war['c'][2] is None
    assert war['c'][1] == [1, 550] and war['t'] == [1, 550]
    lig = _node(out['tree'], 'Лиговский')
    assert lig['c'][0] == [2, 1100] and lig['c'][2] == [2, 500] and lig['t'][0] == 4
    assert out['tree']['c'] == [[2, 1100], [1, 550], [2, 500]] and out['tree']['t'] == [5, 2150]


def test_pivot_empty_report_has_no_false_warnings():
    _fresh(facts=[])
    out = _report(rows=['Store.Name'], columns=['DishGroup.TopParent'],
                  measures=['UniqOrderId'], format='pivot')
    assert out['leaf_count'] == 0 and out['columns'] == [] and not out['warnings']
    assert out['tree']['t'] == [None]


def test_pivot_without_summary_leaves_dash_and_warns():
    _fresh()
    spec = _spec(rows=['Store.Name'], measures=['DishDiscountSumInt', 'UniqOrderId'])
    cat = olap_catalog.get_catalog('SALES')
    data = FakeSession().post_olap(olap_constructor.olap_body(spec, ['Store.Name'], False))
    out = olap_constructor.build_pivot(spec, cat, {'main': dict(data, summary=[])})
    assert out['tree']['t'][0] == 2150 and out['tree']['t'][1] is None
    assert any('Чеков' in w for w in out['warnings'])


def test_pivot_limits():
    _fresh()
    spec = _spec(rows=['DishName'])
    cat = olap_catalog.get_catalog('SALES')
    data = FakeSession().post_olap(olap_constructor.olap_body(spec, ['DishName'], False))
    error = _expect_error(olap_constructor.build_pivot, spec, cat, {'main': data}, 2,
                          code='too_many_rows')
    assert error.extra['limit'] == 2


def test_run_report_uses_cache_for_same_request():
    _fresh()
    payload = dict(rows=['Store.Name'], measures=['DishDiscountSumInt'], format='pivot')
    first = _report(**payload)
    second = _report(**payload)
    assert len(_posts()) == 1
    assert first['meta']['cached'] is False and second['meta']['cached'] is True


def test_cache_never_mixes_answers_of_different_trips_to_iiko():
    """Запрос 'rows' отчёта со столбцами совпадает с отчётом «только строки». Раньше он
    брался из кэша того отчёта (10 минут назад), а ячейки — свежие: итог «ЕДА» по всем
    барам выходил меньше суммы по барам. Кэш — по серии целиком."""
    facts = [dict(fact) for fact in FACTS]
    _fresh(facts=facts)
    first = _report(rows=['DishGroup.TopParent'], measures=['UniqOrderId'], format='pivot')
    assert _node(first['tree'], 'ЕДА')['t'] == [2]                   # o1 и o3
    facts.append(dict(facts[1], **{'UniqOrderId.Id': 'o7', 'Store.Name': 'Варшавская',
                                   'OpenDate.Typed': '2026-09-23'}))   # новая продажа
    payload = dict(rows=['DishGroup.TopParent'], columns=['Store.Name'],
                   measures=['UniqOrderId'], format='pivot')
    second = _report(**payload)
    food = _node(second['tree'], 'ЕДА')
    assert food['c'] == [[2], [1]] and food['t'] == [3], food
    assert second['meta']['cached'] is False and len(_posts()) == 3  # серия — целиком заново
    third = _report(**payload)
    assert third['meta']['cached'] is True and len(_posts()) == 3
    assert olap_client.cache_stats()['waiting'] == 0                 # замки single-flight убраны


def test_cache_memory_limits():
    _fresh()
    saved = olap_client.CACHE_ENTRY_MAX_CELLS, olap_client.CACHE_MAX_CELLS
    try:
        olap_client.CACHE_ENTRY_MAX_CELLS = 3
        _report(rows=['Store.Name'], measures=['DishDiscountSumInt'])      # 2 строки × 2 поля
        assert olap_client.cache_stats()['entries'] == 0                   # крупнее — мимо кэша
        olap_client.CACHE_ENTRY_MAX_CELLS, olap_client.CACHE_MAX_CELLS = 1000, 20
        _report(rows=['Store.Name'], measures=['DishDiscountSumInt'])             # 4 значения
        _report(rows=['DishGroup.TopParent'], measures=['DishDiscountSumInt'])    # 8
        _report(rows=['DishName'], measures=['DishDiscountSumInt'])               # 10
        stats = olap_client.cache_stats()
        assert stats['entries'] == 2 and stats['cells'] == 18, stats          # старейшая ушла
    finally:
        olap_client.CACHE_ENTRY_MAX_CELLS, olap_client.CACHE_MAX_CELLS = saved
        olap_client.clear_cache()


# ------------------------------------------------------------------ список (flat)

def test_flat_sort_limit_having_totals_and_labels():
    _fresh()
    out = _report(rows=['DishGroup.TopParent', 'Store.Name'],
                  measures=['DishDiscountSumInt', 'UniqOrderId'],
                  sort='DishDiscountSumInt', limit=2,
                  having=[{'measure': 'DishDiscountSumInt', 'op': '>=', 'value': 300}])
    assert out['format'] == 'flat'
    assert out['rows'] == [['Напитки Розлив', 'Лиговский', 1250, 3], ['ЕДА', 'Варшавская', 300, 1]] or \
        out['rows'] == [['Напитки Розлив', 'Лиговский', 1250, 3], ['ЕДА', 'Лиговский', 300, 1]]
    assert out['totals'] == {'DishDiscountSumInt': 2150, 'UniqOrderId': 5}
    assert out['row_count'] == 5 and out['matched'] == 3 and out['returned'] == 2 and out['truncated']
    _fresh()
    deleted = _report(rows=['DeletedWithWriteoff'], measures=['DishDiscountSumInt'],
                      include_deleted=True, sort='DeletedWithWriteoff')
    labels = [row[0] for row in deleted['rows']]
    assert labels == ['Не удалено', 'Удалено со списанием']


def test_flat_options_are_checked_before_iiko():
    _fresh()
    base = {'rows': ['Store.Name'], 'measures': ['DishDiscountSumInt']}
    _expect_error(lambda: _report(sort='DishName', **base), text='sort')
    _expect_error(lambda: _report(limit=0, **base), text='limit')
    _expect_error(lambda: _report(having=[{'measure': 'UniqOrderId', 'op': '>', 'value': 1}], **base),
                  text='having')
    assert not _posts()


def test_sort_key_survives_unicode_digits_and_puts_empties_last():
    # «²» — цифра для str.isdigit, но int() её не читает: отчёт падал с 400.
    ordered = sorted(['10²', '2', 'x①', None, '', '1'], key=olap_constructor.sort_key)
    assert ordered.index('1') < ordered.index('2') < ordered.index('10²')
    assert ordered[-2:] == ['', None]


def test_blank_string_and_null_are_different_values():
    """iiko отдаёт и "", и null. Фильтр по null строки с "" не ловит — значит, это разные
    значения: разные подписи, обе в конце списка, в фильтре — отдельными галочками."""
    facts = [dict(fact) for fact in FACTS]
    facts[6]['DishGroup.TopParent'] = ''                     # o5 Stout — пустая строка
    _fresh(facts=facts)
    out = _report(rows=['DishGroup.TopParent'], measures=['DishDiscountSumInt'], format='pivot')
    assert [child['k'] for child in out['tree']['ch']][-2:] == ['', None]
    values = olap_constructor.field_values('SALES', 'DishGroup.TopParent', SALES_BASE)
    assert values['has_empty'] is True
    assert values['values'][-1] == {'value': '', 'label': '(пустая строка)'}
    flat = _report(rows=['DishGroup.TopParent'], measures=['DishDiscountSumInt'],
                   sort='DishGroup.TopParent', order='desc')
    labels = [row[0] for row in flat['rows']]
    assert labels[:2] == ['Напитки Фасовка', 'Напитки Розлив'] and labels[-2:] == ['', None]


# ------------------------------------------------------------------ значения и пресеты

def test_field_values_enum_union_and_empty():
    _fresh()
    out = olap_constructor.field_values('SALES', 'DeletedWithWriteoff', SALES_BASE)
    assert [v['value'] for v in out['values']] == \
        ['NOT_DELETED', 'DELETED_WITHOUT_WRITEOFF', 'DELETED_WITH_WRITEOFF']
    assert out['values'][0]['label'] == 'Не удалено'
    groups = olap_constructor.field_values('SALES', 'DishGroup.TopParent', SALES_BASE, q='напит')
    assert [v['value'] for v in groups['values']] == ['Напитки Розлив', 'Напитки Фасовка']
    assert groups['has_empty'] is True
    body = _posts()[-1]
    assert set(body['filters']) == {'OpenDate.Typed'}              # без «не удалено»
    _expect_error(olap_constructor.field_values, 'SALES', 'DishDiscountSumInt', SALES_BASE,
                  text='показатель')


def test_iiko_presets_conversion():
    _fresh()
    data = olap_constructor.iiko_presets()
    presets = {p['id']: p for p in data['presets']}
    first = presets['p-1']
    assert first['config']['period'] == {'preset': 'last_month'}
    assert first['config']['columns'] == ['DishGroup.TopParent']
    assert {'field': 'SessionNum', 'op': 'range', 'from': 1, 'to': 9, 'include_low': True,
            'include_high': True} in first['config']['filters']
    assert any('Weird' in note for note in first['notes'])
    assert first['config']['include_deleted'] is True
    assert presets['p-2']['config']['period'] == {'from': '2025-03-01', 'to': '2025-03-31'}
    assert presets['p-3']['supported'] is True
    assert presets['p-4']['supported'] is False
    olap_constructor.iiko_presets()
    assert len([e for e in FakeSession.log if e[0] == 'GET' and e[1].endswith('/presets')]) == 1


# ------------------------------------------------------------------ сохранённые отчёты

def test_saved_reports_crud():
    folder = tempfile.mkdtemp()
    olap_saved_reports.set_path(os.path.join(folder, 'explorer_reports.json'))
    try:
        config = {'report_type': 'sales', 'rows': ['Store.Name'], 'measures': ['DishDiscountSumInt'],
                  'filters': [{'field': 'Store.Name', 'op': 'in', 'values': ['Лиговский']}],
                  'period': {'preset': 'last_month'}}
        first = olap_saved_reports.save_report('  Выручка   по барам ', config, 'владелец')
        assert first['name'] == 'Выручка по барам' and first['config']['report_type'] == 'SALES'
        second = olap_saved_reports.save_report('Второй', dict(config, period=None), 'агент')
        updated = olap_saved_reports.save_report('Переименован', config, 'владелец',
                                                 report_id=first['id'])
        assert updated['id'] == first['id'] and updated['created_by'] == 'владелец'
        names = [r['name'] for r in olap_saved_reports.list_reports()]
        assert set(names) == {'Переименован', 'Второй'}
        olap_saved_reports.delete_report(second['id'])
        assert [r['id'] for r in olap_saved_reports.list_reports()] == [first['id']]
        for bad, text in [
            ({'report_type': 'X', 'rows': ['a'], 'measures': ['b']}, 'report_type'),
            ({'report_type': 'SALES', 'rows': [], 'measures': ['b']}, 'пустой отчёт'),
            ({'report_type': 'SALES', 'rows': ['a'], 'measures': []}, 'нет показателей'),
            ({'report_type': 'SALES', 'rows': ['a'], 'measures': ['b'],
              'period': {'preset': 'week'}}, 'preset'),
            ({'report_type': 'SALES', 'rows': ['a'], 'measures': ['b'],
              'filters': [{'field': 'a', 'op': 'eq'}]}, 'op'),
        ]:
            _expect_error(olap_saved_reports.save_report, 'x', bad, 'владелец', text=text)
        _expect_error(olap_saved_reports.save_report, '', config, 'владелец', text='пустым')
        _expect_error(olap_saved_reports.delete_report, 'nope', code='not_found')
        _expect_error(olap_saved_reports.save_report, 'x', config, 'в', 'nope', code='not_found')
        with open(os.path.join(folder, 'explorer_reports.json'), 'w', encoding='utf-8') as handle:
            handle.write('{broken')
        _expect_error(olap_saved_reports.list_reports, text='не читается')
    finally:
        olap_saved_reports.set_path(None)


def test_saved_reports_defaults_dates_ids_and_busy_lock():
    import portalocker
    from contextlib import contextmanager
    folder = tempfile.mkdtemp()
    olap_saved_reports.set_path(os.path.join(folder, 'explorer_reports.json'))
    try:
        base = {'report_type': 'SALES', 'rows': ['Store.Name'], 'measures': ['DishDiscountSumInt']}
        # Без флага — как у построения отчёта: «не удалено» ставится автоматически.
        assert olap_saved_reports.save_report('Агент', base, 'агент')['config']['include_deleted'] is False
        _expect_error(olap_saved_reports.save_report, 'x',
                      dict(base, period={'from': '2026-13-45', 'to': '2026-13-46'}), 'в',
                      text='YYYY-MM-DD')
        _expect_error(olap_saved_reports.save_report, 'x', base, 'в', 123, text='id')
        _expect_error(olap_saved_reports.delete_report, 7, text='id')

        @contextmanager
        def busy(_path):
            raise portalocker.exceptions.AlreadyLocked('занято')
            yield  # noqa: unreachable — форма контекстного менеджера

        original = olap_saved_reports.file_lock
        olap_saved_reports.file_lock = busy
        try:
            error = _expect_error(olap_saved_reports.save_report, 'x', base, 'в', code='busy')
            assert error.status == 503
        finally:
            olap_saved_reports.file_lock = original
    finally:
        olap_saved_reports.set_path(None)


# ------------------------------------------------------------------ Excel

def test_export_workbook_sheets_and_values():
    from openpyxl import load_workbook
    _fresh()
    content, name = olap_export.export_report(dict(
        SALES_BASE, rows=['DishGroup.TopParent', 'DishName'], columns=['Store.Name'],
        measures=['DishDiscountSumInt', 'UniqOrderId']))
    assert name == 'olap_sales_2026-09-21_2026-09-23.xlsx'
    book = load_workbook(io.BytesIO(content))
    assert book.sheetnames == ['Отчёт', 'Данные', 'Параметры']
    report = list(book['Отчёт'].iter_rows(values_only=True))
    assert report[0][2] == 'Варшавская' and report[0][6] == 'Итого'
    assert report[-1][0] == 'Итого' and report[-1][-2:] == (2150, 5)
    group = next(row for row in report if row[0] == 'ЕДА' and row[1] is None)
    assert group[-2:] == (600, 2)
    data = list(book['Данные'].iter_rows(values_only=True))
    assert data[0] == ('Группа блюда 1-го уровня', 'Блюдо', 'Со склада', 'Сумма со скидкой', 'Чеков')
    assert len(data) == 7
    params = {row[0]: row[1] for row in book['Параметры'].iter_rows(values_only=True) if row[0]}
    assert 'Чеков' in params['Итоги из iiko (buildSummary)']
    assert any(key.startswith('Запрос iiko') for key in params)


def test_export_keeps_text_from_iiko_as_text():
    """Название «=1+2» не становится формулой Excel, управляющий символ не роняет выгрузку."""
    from openpyxl import load_workbook
    facts = [dict(fact) for fact in FACTS]
    facts[0]['DishName'] = facts[2]['DishName'] = '=1+2'
    facts[3]['DishName'] = 'Bottle\x0b'
    _fresh(facts=facts)
    content, _name = olap_export.export_report(dict(
        SALES_BASE, rows=['DishName'], measures=['DishDiscountSumInt']))
    book = load_workbook(io.BytesIO(content))
    for sheet in ('Отчёт', 'Данные'):
        cells = [cell for row in book[sheet].iter_rows() for cell in row]
        formulas = [cell for cell in cells if cell.value == '=1+2']
        assert formulas and all(cell.data_type == 's' for cell in formulas), sheet
        assert any(cell.value == 'Bottle' for cell in cells), sheet


# ------------------------------------------------------------------ клиент iiko

class _Response:
    def __init__(self, status, text='', payload=None):
        self.status_code = status
        self.text = text
        self._payload = payload

    def json(self):
        if self._payload is None:
            raise ValueError('not json')
        return self._payload


def _session_with(response_or_error):
    class _Api:
        base_url = 'https://iiko.test/resto/api'

    class _Olap:
        api = _Api()
        token = 'token'

    session = olap_client.IikoSession.__new__(olap_client.IikoSession)
    session._olap = _Olap()
    session._open = True
    original = olap_client.requests.request

    def fake(*args, **kwargs):
        if isinstance(response_or_error, Exception):
            raise response_or_error
        return response_or_error
    olap_client.requests.request = fake
    return session, original


def test_client_turns_iiko_answers_into_readable_errors():
    cases = [
        (_Response(400, 'Filtering is not allowed for field DiscountSum'), 400, 'Filtering is not allowed'),
        (_Response(403, 'License enhancement is required'), 502, 'отказал в доступе'),
        (_Response(500, '<html><body>error</body></html>'), 502, 'HTML'),
        (requests.exceptions.ReadTimeout(), 504, 'не ответил'),
    ]
    for answer, status, text in cases:
        session, original = _session_with(answer)
        try:
            session.post_olap({'reportType': 'SALES'})
        except IikoError as error:
            assert error.status == status and text in error.message, (status, error.message)
        else:
            raise AssertionError('ожидалась IikoError')
        finally:
            olap_client.requests.request = original
    session, original = _session_with(_Response(200, payload={'data': [{'a': 1}], 'summary': []}))
    try:
        assert session.post_olap({'reportType': 'SALES'})['data'] == [{'a': 1}]
    finally:
        olap_client.requests.request = original
    session, original = _session_with(_Response(200, payload={'no': 'data'}))
    try:
        try:
            session.post_olap({})
        except IikoError as error:
            assert 'data' in error.message
        else:
            raise AssertionError('ответ без data — ошибка')
    finally:
        olap_client.requests.request = original


def test_client_retries_only_when_iiko_did_not_start_answering():
    """Таймаут посреди тела ответа requests поднимает как ConnectionError — повтор отправил
    бы тяжёлый отчёт в iiko второй раз. Повтор — только когда связь не установилась."""
    from urllib3.exceptions import ReadTimeoutError
    cases = [
        (requests.exceptions.ConnectionError(ReadTimeoutError(None, None, 'Read timed out.')),
         504, 'не ответил', 1),
        (requests.exceptions.ChunkedEncodingError('Connection broken'), 502, 'оборвал', 1),
        (requests.exceptions.InvalidURL('bad'), 502, 'не выполнен', 1),
        (requests.exceptions.ConnectionError('Connection refused'), 502, 'Нет связи', 2),
    ]
    sleep = olap_client.time.sleep
    olap_client.time.sleep = lambda seconds: None
    try:
        for error, status, text, attempts in cases:
            session, original = _session_with(error)
            raising = olap_client.requests.request
            calls = []

            def counted(*args, **kwargs):
                calls.append(1)
                return raising(*args, **kwargs)
            olap_client.requests.request = counted
            try:
                session.post_olap({'reportType': 'SALES'})
            except IikoError as caught:
                assert caught.status == status and text in caught.message, caught.message
                assert len(calls) == attempts, (type(error).__name__, len(calls))
            else:
                raise AssertionError('ожидалась IikoError')
            finally:
                olap_client.requests.request = original
    finally:
        olap_client.time.sleep = sleep


def test_session_holds_slots_from_login_to_logout():
    slots = olap_client._slots

    class FailingLogin:
        def __enter__(self):
            raise IikoError('нет входа', 502)

        def __exit__(self, *exc):
            return False

    olap_client.set_session_factory(lambda: FailingLogin())
    try:
        with olap_client.open_session(parallel=2):
            raise AssertionError('вход не должен был пройти')
    except IikoError as error:
        assert error.message == 'нет входа'
    assert slots.acquire(2, 0)                   # места вернулись после неудачного входа
    saved = olap_client.WAIT_SLOT_S
    olap_client.WAIT_SLOT_S = 0.05
    olap_client.set_session_factory(lambda: FakeSession())
    try:
        try:
            olap_client.open_session().__enter__()
        except IikoError as error:
            assert error.status == 503           # все места заняты — честный отказ
        else:
            raise AssertionError('сессия без свободного места')
    finally:
        olap_client.WAIT_SLOT_S = saved
        slots.release(2)
    assert slots.acquire(2, 0)
    slots.release(2)


# ------------------------------------------------------------------ маршруты

_APP = None


def _client():
    global _APP
    if _APP is None:
        from flask import Flask
        from routes.explorer import explorer_bp
        app = Flask(__name__, template_folder=os.path.join(os.path.dirname(HERE), 'templates'))
        app.jinja_env.globals['app_version'] = 'test'
        app.secret_key = 'test'
        app.register_blueprint(explorer_bp)
        _APP = app
    return _APP.test_client()


def test_routes_report_values_columns_presets():
    _fresh()
    client = _client()
    res = client.get('/api/explorer/columns?report_type=TRANSACTIONS&compact=1')
    assert res.status_code == 200 and res.get_json()['field_count'] == 116
    assert client.get('/api/explorer/columns?report_type=OLAP').status_code == 400
    body = dict(SALES_BASE, rows=['Store.Name'], measures=['DishDiscountSumInt', 'UniqOrderId'])
    res = client.post('/api/explorer/report', json=dict(body, format='pivot'))
    assert res.status_code == 200 and res.get_json()['tree']['t'] == [2150, 5]
    res = client.post('/api/explorer/report', json=body)
    assert res.get_json()['format'] == 'flat'
    res = client.post('/api/explorer/report', json=dict(body, colour='red'))
    assert res.status_code == 400 and 'colour' in res.get_json()['error']
    res = client.post('/api/explorer/report', json=dict(body, rows=['Nope']))
    assert res.status_code == 400 and res.get_json()['code'] == 'unknown_field'
    res = client.get('/api/explorer/values?report_type=SALES&field=Store.Name&period=last_week')
    assert res.status_code == 200
    assert client.get('/api/explorer/values?report_type=SALES').status_code == 400
    assert client.get('/api/explorer/values?report_type=SALES&field=Store.Name&period=last_week'
                      '&limit=0').status_code == 400
    res = client.get('/api/explorer/presets')
    assert res.status_code == 200 and len(res.get_json()['presets']) == 4


def test_routes_pass_iiko_error_text():
    _fresh(olap_error=IikoError('iiko отклонил запрос: Filtering is not allowed.', 400))
    res = _client().post('/api/explorer/report', json=dict(
        SALES_BASE, rows=['Store.Name'], measures=['DishDiscountSumInt']))
    assert res.status_code == 400
    assert res.get_json() == {'error': 'iiko отклонил запрос: Filtering is not allowed.', 'code': 'iiko'}


def test_routes_saved_and_export_and_page():
    _fresh()
    folder = tempfile.mkdtemp()
    olap_saved_reports.set_path(os.path.join(folder, 'explorer_reports.json'))
    try:
        client = _client()
        assert client.get('/api/explorer/saved').get_json() == {'reports': []}
        config = {'report_type': 'SALES', 'rows': ['Store.Name'], 'measures': ['DishDiscountSumInt'],
                  'period': {'preset': 'last_week'}}
        res = client.post('/api/explorer/saved', json={'name': 'Бары', 'config': config})
        assert res.status_code == 200
        report_id = res.get_json()['report']['id']
        assert res.get_json()['report']['created_by'] == 'неизвестно'
        assert client.post('/api/explorer/saved', json={'name': 'x'}).status_code == 400
        res = client.post('/api/explorer/saved', json={'name': 'x', 'config': config, 'id': 5})
        assert res.status_code == 400 and 'id' in res.get_json()['error']
        assert client.delete('/api/explorer/saved/' + report_id).status_code == 200
        assert client.delete('/api/explorer/saved/' + report_id).status_code == 404
    finally:
        olap_saved_reports.set_path(None)
    res = client.post('/api/explorer/export', json=dict(
        SALES_BASE, rows=['Store.Name'], measures=['DishDiscountSumInt']))
    assert res.status_code == 200 and res.mimetype.endswith('spreadsheetml.sheet')
    assert 'olap_sales_2026-09-21_2026-09-23.xlsx' in res.headers['Content-Disposition']
    page = client.get('/explorer')
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    assert 'id="exBoot"' in html and 'explorer/page.js' in html
    boot = json.loads(html.split('id="exBoot">', 1)[1].split('</script>', 1)[0])
    assert boot['default']['report_type'] == 'SALES'
    assert [p['id'] for p in boot['periods']] == list(olap_constructor.PERIOD_PRESETS)


def teardown_module(module):
    olap_client.set_session_factory(None)
    olap_client.clear_cache()
    olap_catalog.reset_cache()
    olap_constructor.reset_presets_cache()


if __name__ == '__main__':
    tests = [(name, fn) for name, fn in sorted(globals().items())
             if name.startswith('test_') and callable(fn)]
    failed = 0
    for name, fn in tests:
        try:
            fn()
        except Exception as error:  # noqa: BLE001 — самозапуск печатает все падения
            failed += 1
            print('FAIL ' + name + ': ' + type(error).__name__ + ': ' + str(error)[:600])
        else:
            print('ok   ' + name)
    teardown_module(None)
    print(str(len(tests) - failed) + '/' + str(len(tests)) + ' passed')
    sys.exit(1 if failed else 0)
